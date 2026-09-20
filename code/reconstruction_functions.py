from __future__ import annotations

import json
import hashlib
import math
import re
import argparse
from datetime import timezone
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import openpyxl
import pandas as pd
from sklearn.base import clone
from sklearn.compose import TransformedTargetRegressor
from sklearn.ensemble import (
    ExtraTreesRegressor,
    GradientBoostingRegressor,
    HistGradientBoostingRegressor,
    RandomForestRegressor,
)
from sklearn.impute import SimpleImputer
from sklearn.linear_model import HuberRegressor, Ridge
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import RobustScaler, StandardScaler


HERE = Path(__file__).resolve().parent
WORKSPACE = HERE.parents[1]
SOURCE = HERE.parent / "data" / "Farm-scale_Biodigester.xlsx"

SHEETS = {
    "RI-FLEX": "RI-FLEX",
    "R2-FLEX": "R2-FLEX",
    "R3-FIXED DOME": "R3",
    "R4-FIXED DOME": "R4",
}

METER_SWAP_DATES = {
    "RI-FLEX": {date(2024, 10, 8)},
    "R2-FLEX": {date(2024, 10, 8)},
    "R3": {date(2024, 3, 8), date(2024, 10, 8)},
    "R4": {date(2024, 3, 8), date(2024, 10, 8)},
}

# Source-supported date recovery. R3 row 29 has no date, but its counter value
# (113,601) equals R4's 7 March counter and precedes R3's dated 8 March row.
DATE_OVERRIDES = {("R3-FIXED DOME", 29): date(2024, 3, 8)}



def json_default(value: Any) -> Any:
    if isinstance(value, (date, datetime, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, np.ndarray):
        return value.tolist()
    raise TypeError(type(value).__name__)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def number(value: Any) -> float:
    if value is None:
        return float("nan")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    match = re.search(r"[-+]?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group()) if match else float("nan")


def date_candidates(raw: Any) -> list[date]:
    candidates: list[date] = []
    if isinstance(raw, datetime):
        base = raw.date()
        candidates.append(base)
        if base.day <= 12 and base.month <= 12 and base.day != base.month:
            candidates.append(date(base.year, base.day, base.month))
    elif isinstance(raw, date):
        candidates.append(raw)
    elif isinstance(raw, str):
        match = re.search(r"(\d{1,2})\D+(\d{1,2})\D+(\d{4})", raw.strip())
        if match:
            first, second, year = map(int, match.groups())
            for month, day in ((second, first), (first, second)):
                try:
                    candidates.append(date(year, month, day))
                except ValueError:
                    pass
    return sorted(set(candidates))


def transition_cost(previous: date, current: date) -> float:
    gap = (current - previous).days
    if gap < 0:
        return float("inf")
    if gap == 0:
        return 4.0
    if gap <= 3:
        return 0.05 * (gap - 1)
    if gap <= 7:
        return 0.4 * gap
    return 3.0 * gap


def choose_monotone_dates(rows: list[tuple[int, Any]]) -> dict[int, date]:
    """Resolve day/month ambiguity from source-row chronology with a transparent DP."""
    candidate_map = {row: date_candidates(raw) for row, raw in rows}
    usable = [(row, candidate_map[row]) for row, _ in rows if candidate_map[row]]
    if not usable:
        return {}

    costs: list[dict[date, float]] = []
    back: list[dict[date, date | None]] = []
    for i, (_, candidates) in enumerate(usable):
        current_costs: dict[date, float] = {}
        current_back: dict[date, date | None] = {}
        for candidate in candidates:
            if i == 0:
                # The sheets begin in January 2024; this only breaks the first-row ambiguity.
                initial = abs((candidate - date(2024, 1, 1)).days) * 0.02
                current_costs[candidate] = initial
                current_back[candidate] = None
                continue
            best_cost = float("inf")
            best_prev: date | None = None
            for previous, previous_cost in costs[-1].items():
                trial = previous_cost + transition_cost(previous, candidate)
                if trial < best_cost:
                    best_cost = trial
                    best_prev = previous
            current_costs[candidate] = best_cost
            current_back[candidate] = best_prev
        costs.append(current_costs)
        back.append(current_back)

    last = min(costs[-1], key=costs[-1].get)
    selected: list[date] = [last]
    for i in range(len(usable) - 1, 0, -1):
        previous = back[i][selected[-1]]
        if previous is None:
            raise RuntimeError("Date-path backtracking failed")
        selected.append(previous)
    selected.reverse()
    return {row: chosen for (row, _), chosen in zip(usable, selected)}


def raw_date_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    return "" if value is None else str(value)


def parse_sheet(ws: openpyxl.worksheet.worksheet.Worksheet, reactor: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = [
        (row, ws.cell(row, 1).value)
        for row in range(2, ws.max_row + 1)
        if ws.cell(row, 1).value is not None or (ws.title, row) in DATE_OVERRIDES
    ]
    selected_dates = choose_monotone_dates(rows)
    for (sheet_name, row), recovered_date in DATE_OVERRIDES.items():
        if ws.title == sheet_name:
            selected_dates[row] = recovered_date
    records: list[dict[str, Any]] = []
    date_audit: list[dict[str, Any]] = []
    for row, raw in rows:
        parsed = selected_dates.get(row)
        if parsed is None:
            continue
        original_candidates = date_candidates(raw)
        if isinstance(raw, datetime):
            original = raw.date()
        elif isinstance(raw, date):
            original = raw
        else:
            original = original_candidates[0] if original_candidates else None
        overridden = (ws.title, row) in DATE_OVERRIDES
        corrected = overridden or (original is not None and parsed != original)
        date_audit.append(
            {
                "reactor": reactor,
                "source_sheet": ws.title,
                "source_row": row,
                "raw_date": raw_date_text(raw),
                "selected_date": parsed.isoformat(),
                "candidate_dates": "|".join(item.isoformat() for item in original_candidates),
                "date_corrected": corrected,
                "reason": (
                    "source-supported undated meter-swap baseline recovery"
                    if overridden
                    else "source-order day/month disambiguation"
                    if corrected
                    else "as recorded/unambiguous"
                ),
            }
        )
        records.append(
            {
                "reactor": reactor,
                "source_sheet": ws.title,
                "source_row": row,
                "date": parsed,
                "meter_reading": number(ws.cell(row, 7).value),
                "air_temp_C": number(ws.cell(row, 8).value),
                "digester_temp_C": number(ws.cell(row, 9).value),
                "manure_kg": number(ws.cell(row, 5).value),
                "water_kg": number(ws.cell(row, 6).value),
                "comment": str(ws.cell(row, 28).value or "").strip(),
            }
        )
    return pd.DataFrame(records), pd.DataFrame(date_audit)


def build_source_table() -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = openpyxl.load_workbook(SOURCE, data_only=False, read_only=False)
    tables: list[pd.DataFrame] = []
    audits: list[pd.DataFrame] = []
    for sheet_name, reactor in SHEETS.items():
        table, audit = parse_sheet(workbook[sheet_name], reactor)
        tables.append(table)
        audits.append(audit)
    source = pd.concat(tables, ignore_index=True)
    date_audit = pd.concat(audits, ignore_index=True)
    return source, date_audit


def collapse_duplicate_dates(source: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    source = source.sort_values(["reactor", "date", "source_row"]).copy()
    source["duplicate_date_count"] = source.groupby(["reactor", "date"])["source_row"].transform("size")
    duplicates = source[source["duplicate_date_count"] > 1].copy()
    # Last source row is the end-of-day counter used as the next baseline; every duplicate remains in the audit.
    collapsed = source.groupby(["reactor", "date"], as_index=False, sort=True).tail(1).copy()
    return collapsed.sort_values(["reactor", "date"]), duplicates


def build_target_ledger(source: pd.DataFrame) -> pd.DataFrame:
    records: list[dict[str, Any]] = []
    for reactor, group in source.groupby("reactor", sort=True):
        group = group.sort_values(["date", "source_row"])
        last_reading: float | None = None
        last_date: date | None = None
        segment = 0
        for row in group.itertuples(index=False):
            current = row._asdict()
            current_date: date = current["date"]
            reading = current["meter_reading"]
            temp = current["air_temp_C"]
            reason = "eligible"
            eligible = True
            elapsed = float("nan")
            increment = float("nan")
            target = float("nan")

            if current_date in METER_SWAP_DATES[reactor]:
                segment += 1
                eligible = False
                reason = f"source-supported meter-swap boundary {current_date.isoformat()}; baseline reset"
                if np.isfinite(reading):
                    last_reading = float(reading)
                    last_date = current_date
            elif not np.isfinite(reading):
                eligible = False
                reason = "missing cumulative meter reading; baseline unchanged"
            elif last_reading is None or last_date is None:
                eligible = False
                reason = "segment baseline only"
                last_reading = float(reading)
                last_date = current_date
            else:
                elapsed = float((current_date - last_date).days)
                increment = float(reading - last_reading)
                if elapsed <= 0:
                    eligible = False
                    reason = "non-positive elapsed time; baseline unchanged"
                elif increment < 0:
                    eligible = False
                    reason = "cumulative counter reversal; reading rejected and last accepted baseline retained"
                elif not np.isfinite(temp) or float(temp) <= 0.0:
                    eligible = False
                    reason = "missing/non-positive temperature for STP correction; valid counter becomes next baseline"
                    last_reading = float(reading)
                    last_date = current_date
                else:
                    target = increment / elapsed * 273.0 / (273.0 + float(temp))
                    last_reading = float(reading)
                    last_date = current_date

            current.update(
                {
                    "meter_segment": segment,
                    "elapsed_days": elapsed,
                    "meter_increment": increment,
                    "target_stp_mL_day": target,
                    "eligible_target": eligible and np.isfinite(target),
                    "eligibility_reason": reason,
                }
            )
            records.append(current)
    return pd.DataFrame(records).sort_values(["date", "reactor", "source_row"]).reset_index(drop=True)

